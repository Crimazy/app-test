"""Continuous RC beam analysis & design (Singly / Doubly reinforced, SDM).

Port of the workbook "Beam_Design_Singly_Doubly.xlsm" (idea based on the
SUTStructor beam module):

  * Analysis  : direct stiffness method, 2 DOF/node (v, theta), supports
                Pin / Fixed / Free / Column (rotational spring 4EI/H of the
                columns above and below, far ends fixed), moment release at
                nodes ("simple beams in series").
  * Loads     : DL on all spans, LL on each span separately -> exact
                pattern-LL envelope (sum of +ve / -ve parts) and 1.4DL.
  * Design    : 3 locations per span (left / mid / right), automatic singly /
                doubly reinforcement, strain-compatibility capacity check of
                the bars provided, stirrups.

Units follow the workbook: t, m (analysis); cm, cm^2, ksc (design);
moments t-m, shears t.  Moment sign: + sagging (bottom in tension).
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field, replace

NDIV = 40
MAXS = 10

CODE99 = "ACI 318-99 / วสท. (1.4D+1.7L)"
CODE19 = "ACI 318-19 / มยผ. (1.2D+1.6L)"
CODES = [CODE99, CODE19]
SUPPORT_TYPES = ["Pin", "Fixed", "Free", "Column"]
POSITIONS = ["ซ้าย", "กลาง", "ขวา"]
STIRRUP_SIZES = [6, 9, 10, 12]
BAR_SIZES = [12, 16, 20, 25, 28, 32]


# --------------------------------------------------------------------------
# Input
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Settings:
    fc: float = 280.0          # ksc
    fy: float = 4000.0         # ksc
    fys: float = 2400.0        # ksc
    Es: float = 2_040_000.0    # ksc
    gc: float = 2.4            # t/m^3
    code: str = CODE99
    pattern: bool = True
    self_weight: bool = True
    cover: float = 3.0         # cm, to stirrup
    stir: float = 9.0          # mm
    legs: int = 2
    db: float = 20.0           # mm, default main bar
    all_simple: bool = False
    load_factors: tuple | None = None   # (fD, fL); None = follow the code

    @property
    def is99(self) -> bool:
        return self.code.startswith("ACI 318-99")

    @property
    def fD(self) -> float:
        if self.load_factors:
            return self.load_factors[0]
        return 1.4 if self.is99 else 1.2

    @property
    def fL(self) -> float:
        if self.load_factors:
            return self.load_factors[1]
        return 1.7 if self.is99 else 1.6

    @property
    def phib(self) -> float:
        return 0.9

    @property
    def phiv(self) -> float:
        return 0.85 if self.is99 else 0.75

    @property
    def beta1(self) -> float:
        return max(0.65, 0.85 if self.fc <= 280 else 0.85 - 0.05 * (self.fc - 280) / 70)

    @property
    def es_cu(self) -> float:
        """0.003 Es (= 6120 ksc for Es = 2.04e6)."""
        return 0.003 * self.Es

    @property
    def rhob(self) -> float:
        return 0.85 * self.beta1 * self.fc / self.fy * self.es_cu / (self.es_cu + self.fy)

    @property
    def rhomax(self) -> float:
        if self.is99:
            return 0.75 * self.rhob
        return 0.85 * self.beta1 * self.fc / self.fy * 0.003 / (0.003 + 0.005)

    @property
    def rhomin(self) -> float:
        return max(0.8 * math.sqrt(self.fc) / self.fy, 14 / self.fy)

    @property
    def Ec(self) -> float:
        return 15100 * math.sqrt(self.fc)   # ksc


@dataclass(frozen=True)
class Span:
    L: float                   # m, column centre to column centre
    b: float                   # cm
    h: float                   # cm
    name: str = "B1"


@dataclass(frozen=True)
class Node:
    kind: str = "Pin"          # Pin / Fixed / Free / Column
    c: float = 0.0             # column width along the beam (cm)
    lower: tuple = (0.0, 0.0, 0.0)   # column below: b (cm), h (cm), H (m)
    upper: tuple = (0.0, 0.0, 0.0)   # column above
    continuous: bool = True    # False = moment released here


@dataclass(frozen=True)
class Load:
    span: int                  # 1-based
    case: str = "DL"           # DL / LL
    kind: str = "UDL"          # UDL / Point
    w1: float = 0.0            # t/m or P (t)
    w2: float | None = None    # t/m (None = w1)
    a: float | None = None     # m from the left end of the span
    b: float | None = None     # m (None = L)


class InputError(ValueError):
    pass


# --------------------------------------------------------------------------
# Analysis (port of BeamAnalysis.bas)
# --------------------------------------------------------------------------
def _col_k(Ec_tm2: float, col: tuple) -> float:
    b, h, H = col
    b, h = b / 100.0, h / 100.0
    if b > 0 and h > 0 and H > 0:
        return 4.0 * Ec_tm2 * (b * h ** 3 / 12.0) / H
    return 0.0


def _gauss(A, b):
    """Gaussian elimination with partial pivoting (as in the macro)."""
    n = len(b)
    A = [row[:] for row in A]
    b = b[:]
    scl = max(abs(A[i][i]) for i in range(n)) or 1.0
    for k in range(n):
        p = max(range(k, n), key=lambda i: abs(A[i][k]))
        if abs(A[p][k]) < 1e-9 * scl:
            return None
        if p != k:
            A[k], A[p] = A[p], A[k]
            b[k], b[p] = b[p], b[k]
        for i in range(k + 1, n):
            t = A[i][k] / A[k][k]
            if t != 0:
                for j in range(k, n):
                    A[i][j] -= t * A[k][j]
                b[i] -= t * b[k]
    x = [0.0] * n
    for i in range(n - 1, -1, -1):
        t = b[i] - sum(A[i][j] * x[j] for j in range(i + 1, n))
        x[i] = t / A[i][i]
    return x


@dataclass
class _Ld:
    span: int
    case: int                  # 0 = DL, 1 = LL
    kind: int                  # 1 = distributed, 2 = point
    w1: float
    w2: float
    a: float
    b: float


class BeamModel:
    def __init__(self, st: Settings, spans: list[Span], nodes: list[Node], loads: list[Load]):
        self.st, self.spans = st, spans
        nS = len(spans)
        if not 1 <= nS <= MAXS:
            raise InputError("จำนวนช่วงต้องอยู่ระหว่าง 1 – 10")
        if len(nodes) < nS + 1:
            raise InputError(f"ต้องกำหนดจุดรองรับ {nS + 1} จุด")
        if st.fc <= 0:
            raise InputError("fc' ต้องมากกว่า 0")
        self.nS = nS
        self.nodes = nodes[: nS + 1]
        self.Ec = st.Ec * 10.0                 # ksc -> t/m^2
        self.X0 = [0.0]
        self.dshear = []
        for i, s in enumerate(spans, 1):
            if s.L <= 0 or s.b <= 0 or s.h <= 0:
                raise InputError(f"ช่วงที่ {i}: L, b, h ต้องมากกว่า 0")
            self.X0.append(self.X0[-1] + s.L)
            self.dshear.append(s.h - st.cover - st.stir / 10.0 - st.db / 20.0)

        self.kind, self.colw, self.krot, self.brk = [], [], [], []
        for i, nd in enumerate(self.nodes, 1):
            kind = (nd.kind or "Free").strip().upper()
            if kind not in ("PIN", "FIXED", "FREE", "COLUMN"):
                raise InputError(f"จุดที่ {i}: ชนิดจุดรองรับต้องเป็น Pin / Fixed / Free / Column")
            brk = 1 < i <= nS and (st.all_simple or not nd.continuous)
            k = _col_k(self.Ec, nd.lower) + _col_k(self.Ec, nd.upper) if kind == "COLUMN" else 0.0
            self.kind.append(kind)
            self.colw.append(nd.c or 0.0)
            self.krot.append(0.0 if brk else k)
            self.brk.append(brk)
        self.ndof = 2 * (nS + 1)
        self.exdof = []
        for i in range(nS + 1):
            if self.brk[i]:
                self.exdof.append(self.ndof)
                self.ndof += 1
            else:
                self.exdof.append(None)

        self.loads: list[_Ld] = []
        for r, ld in enumerate(loads, 1):
            sp = int(ld.span)
            if sp < 1 or sp > nS:
                raise InputError(f"น้ำหนักแถว {r}: หมายเลขช่วงไม่ถูกต้อง")
            L = spans[sp - 1].L
            case = 1 if str(ld.case).strip().upper() == "LL" else 0
            if str(ld.kind).strip().upper() == "POINT":
                a = L / 2 if ld.a is None else ld.a
                if a < 0 or a > L + 1e-6:
                    raise InputError(f"น้ำหนักแถว {r}: ตำแหน่ง a ของแรงจุดอยู่นอกช่วง")
                self.loads.append(_Ld(sp, case, 2, ld.w1, 0.0, a, 0.0))
            else:
                w2 = ld.w1 if ld.w2 is None else ld.w2
                a = 0.0 if ld.a is None else ld.a
                b = L if ld.b is None else ld.b
                if a < 0 or b > L + 1e-6 or b <= a:
                    raise InputError(f"น้ำหนักแถว {r}: ช่วง a–b ของน้ำหนักแผ่ไม่ถูกต้อง")
                self.loads.append(_Ld(sp, case, 1, ld.w1, w2, a, b))
        if st.self_weight:
            for i, s in enumerate(spans, 1):
                w = st.gc * s.b / 100 * s.h / 100
                self.loads.append(_Ld(i, 0, 1, w, w, 0.0, s.L))

        # element end forces per case c (0 = DL, k = LL on span k)
        self.EV, self.EM, self.ER, self.EMR = {}, {}, {}, {}
        for c in range(nS + 1):
            if not self._solve(c):
                raise InputError("โครงสร้างไม่เสถียร — ตรวจจุดรองรับ (คานเดี่ยวต้องมี 2 จุดรองรับ; "
                                 "คานยื่นต้องต่อเนื่องกับช่วงข้างเคียง)")

    # -- element helpers --------------------------------------------------
    def _in_case(self, ld: _Ld, c: int) -> bool:
        return ld.case == 0 if c == 0 else (ld.case == 1 and ld.span == c)

    def _ke(self, s: int):
        sp = self.spans[s - 1]
        L = sp.L
        k = self.Ec * (sp.b / 100) * (sp.h / 100) ** 3 / 12 / L ** 3
        return [[12 * k, 6 * L * k, -12 * k, 6 * L * k],
                [6 * L * k, 4 * L * L * k, -6 * L * k, 2 * L * L * k],
                [-12 * k, -6 * L * k, 12 * k, -6 * L * k],
                [6 * L * k, 2 * L * L * k, -6 * L * k, 4 * L * L * k]]

    @staticmethod
    def _add_point_fef(L, P, a, fe):
        b = L - a
        fe[0] += P * b * b * (3 * a + b) / L ** 3
        fe[1] += P * a * b * b / L ** 2
        fe[2] += P * a * a * (a + 3 * b) / L ** 3
        fe[3] -= P * a * a * b / L ** 2

    def _fef(self, c: int, s: int):
        fe = [0.0] * 4
        L = self.spans[s - 1].L
        for ld in self.loads:
            if ld.span != s or not self._in_case(ld, c):
                continue
            if ld.kind == 2:
                self._add_point_fef(L, ld.w1, ld.a, fe)
            else:
                n = 200
                dx = (ld.b - ld.a) / n
                for j in range(1, n + 1):
                    x = ld.a + (j - 0.5) * dx
                    w = ld.w1 + (ld.w2 - ld.w1) * (x - ld.a) / (ld.b - ld.a)
                    self._add_point_fef(L, w * dx, x, fe)
        return fe

    def _dof(self, s: int):
        left_rot = self.exdof[s - 1] if self.exdof[s - 1] is not None else 2 * s - 1
        return [2 * s - 2, left_rot, 2 * s, 2 * s + 1]

    def _solve(self, c: int) -> bool:
        nd = self.ndof
        K = [[0.0] * nd for _ in range(nd)]
        F = [0.0] * nd
        for s in range(1, self.nS + 1):
            ke, fe, dof = self._ke(s), self._fef(c, s), self._dof(s)
            for i in range(4):
                F[dof[i]] -= fe[i]
                for j in range(4):
                    K[dof[i]][dof[j]] += ke[i][j]
        for i in range(self.nS + 1):
            K[2 * i + 1][2 * i + 1] += self.krot[i]

        def fix(q):
            for j in range(nd):
                K[q][j] = K[j][q] = 0.0
            K[q][q] = 1.0
            F[q] = 0.0

        for i, kind in enumerate(self.kind):
            if kind in ("PIN", "COLUMN"):
                fix(2 * i)
            elif kind == "FIXED":
                fix(2 * i)
                fix(2 * i + 1)
                if self.exdof[i] is not None:
                    fix(self.exdof[i])
        d = _gauss(K, F)
        if d is None:
            return False
        for s in range(1, self.nS + 1):
            ke, fe, dof = self._ke(s), self._fef(c, s), self._dof(s)
            fo = [fe[i] + sum(ke[i][j] * d[dof[j]] for j in range(4)) for i in range(4)]
            self.EV[c, s], self.EM[c, s], self.ER[c, s], self.EMR[c, s] = fo
        return True

    # -- internal forces ----------------------------------------------------
    @staticmethod
    def _load_left(ld: _Ld, x: float):
        if ld.kind == 2:
            if x > ld.a + 1e-7:
                return ld.w1, ld.w1 * (x - ld.a)
            return 0.0, 0.0
        if x <= ld.a:
            return 0.0, 0.0
        xe = min(x, ld.b)
        we = ld.w1 + (ld.w2 - ld.w1) * (xe - ld.a) / (ld.b - ld.a)
        R = (ld.w1 + we) / 2 * (xe - ld.a)
        if abs(ld.w1 + we) > 1e-9:
            cen = (xe - ld.a) * (ld.w1 + 2 * we) / (3 * (ld.w1 + we))
        else:
            cen = (xe - ld.a) / 2
        return R, R * (x - ld.a - cen)

    def M(self, c: int, s: int, x: float) -> float:
        m = -self.EM[c, s] + self.EV[c, s] * x
        for ld in self.loads:
            if ld.span == s and self._in_case(ld, c):
                m -= self._load_left(ld, x)[1]
        return m

    def V(self, c: int, s: int, x: float) -> float:
        v = self.EV[c, s]
        for ld in self.loads:
            if ld.span == s and self._in_case(ld, c):
                v -= self._load_left(ld, x)[0]
        return v

    def _envelope(self, dl: float, llv: list[float]):
        fD, fL = self.st.fD, self.st.fL
        if self.st.pattern:
            emax = fD * dl + fL * sum(v for v in llv if v > 0)
            emin = fD * dl + fL * sum(v for v in llv if v <= 0)
        else:
            emax = emin = fD * dl + fL * sum(llv)
        e14 = 1.4 * dl
        return max(emax, e14), min(emin, e14)

    def env_M(self, s: int, x: float):
        return self._envelope(self.M(0, s, x), [self.M(k, s, x) for k in range(1, self.nS + 1)])

    def env_V(self, s: int, x: float):
        return self._envelope(self.V(0, s, x), [self.V(k, s, x) for k in range(1, self.nS + 1)])

    # -- results ------------------------------------------------------------
    def span_type(self, s: int) -> str:
        if self.kind[s - 1] == "FREE" or self.kind[s] == "FREE":
            return "Cantilever"
        if (s == 1 or self.brk[s - 1]) and (s == self.nS or self.brk[s]):
            return "Simple"
        return "Continuous"

    def stations(self) -> list[dict]:
        out = []
        for s in range(1, self.nS + 1):
            L = self.spans[s - 1].L
            for j in range(NDIV + 1):
                x = L * j / NDIV
                m1, m2 = self.env_M(s, x)
                v1, v2 = self.env_V(s, x)
                out.append(dict(span=s, x=x, X=self.X0[s - 1] + x, Mmax=m1, Mmin=m2, Vmax=v1, Vmin=v2))
        return out

    def summary(self) -> list[dict]:
        out = []
        for s in range(1, self.nS + 1):
            L = self.spans[s - 1].L
            xl = 0.0 if self.kind[s - 1] == "FREE" else self.colw[s - 1] / 200
            xr = L if self.kind[s] == "FREE" else L - self.colw[s] / 200
            xl, xr = min(xl, L / 2), max(xr, L / 2)
            ml_max, ml_min = self.env_M(s, xl)
            mp_max, mp_x, mn_mid, v_mid = -1e30, 0.0, 1e30, 0.0
            for j in range(NDIV + 1):
                x = L * j / NDIV
                a1, a2 = self.env_M(s, x)
                if a1 > mp_max:
                    mp_max, mp_x = a1, x
                if NDIV / 4 <= j <= 3 * NDIV / 4:
                    mn_mid = min(mn_mid, a2)
                    b1, b2 = self.env_V(s, x)
                    v_mid = max(v_mid, abs(b1), abs(b2))
            for ld in self.loads:   # peaks under point loads
                if ld.span == s and ld.kind == 2:
                    a1, a2 = self.env_M(s, ld.a)
                    if a1 > mp_max:
                        mp_max, mp_x = a1, ld.a
                    if L / 4 <= ld.a <= 3 * L / 4:
                        mn_mid = min(mn_mid, a2)
            mr_max, mr_min = self.env_M(s, xr)
            x = 0.0 if self.kind[s - 1] == "FREE" else min(xl + self.dshear[s - 1] / 100, L / 2)
            vl = max(abs(v) for v in self.env_V(s, x))
            x = L if self.kind[s] == "FREE" else max(xr - self.dshear[s - 1] / 100, L / 2)
            vr = max(abs(v) for v in self.env_V(s, x))
            sp = self.spans[s - 1]
            out.append(dict(span=s, L=L, b=sp.b, h=sp.h, type=self.span_type(s),
                            c_left=self.colw[s - 1], c_right=self.colw[s],
                            ML_max=ml_max, ML_min=ml_min, Mpos_max=mp_max, Mpos_x=mp_x,
                            Mmid_min=mn_mid, MR_max=mr_max, MR_min=mr_min,
                            Vu_left=vl, Vu_mid=v_mid, Vu_right=vr, d=self.dshear[s - 1]))
        return out

    def _node_sum(self, A, B, c, j):
        r = 0.0
        if j <= self.nS:
            r += A[c, j]
        if j > 1:
            r += B[c, j - 1]
        return r

    def reactions(self) -> list[dict]:
        out = []
        for j in range(1, self.nS + 2):
            rD = self._node_sum(self.EV, self.ER, 0, j)
            llv = [self._node_sum(self.EV, self.ER, k, j) for k in range(1, self.nS + 1)]
            rmx, rmn = self._envelope(rD, llv)
            mD = self._node_sum(self.EM, self.EMR, 0, j)
            mll = [self._node_sum(self.EM, self.EMR, k, j) for k in range(1, self.nS + 1)]
            mmx, mmn = self._envelope(mD, mll)
            out.append(dict(node=j, X=self.X0[j - 1], kind=self.kind[j - 1], RD=rD, RL=sum(llv),
                            Ru_max=rmx, Ru_min=rmn, Mu_max=mmx, Mu_min=mmn))
        return out


# --------------------------------------------------------------------------
# Minimum depth (INPUT!H30), span classification as in the workbook
# --------------------------------------------------------------------------
def span_kind_th(nodes: list[Node], s: int, nS: int, all_simple: bool) -> str:
    kL, kR = nodes[s - 1].kind, nodes[s].kind
    if kL == "Free" or kR == "Free":
        return "คานยื่น"
    left_simple = s == 1 or all_simple or not nodes[s - 1].continuous
    right_simple = s == nS or all_simple or not nodes[s].continuous
    return "คานเดี่ยว" if left_simple and right_simple else "คานต่อเนื่อง"


def min_depth(st: Settings, nodes: list[Node], s: int, nS: int, span: Span) -> dict:
    kind = span_kind_th(nodes, s, nS, st.all_simple)
    if kind == "คานยื่น":
        div = 8.0
    else:
        def cont(i, inner):   # end continuous or fixed
            return (inner and not (st.all_simple or not nodes[i].continuous)) or nodes[i].kind == "Fixed"
        l_c = cont(s - 1, s > 1)
        r_c = cont(s, s < nS)
        div = 21.0 if (l_c and r_c) else (18.5 if (l_c or r_c) else 16.0)
    hmin = span.L * 100 / div * (0.4 + st.fy / 7000)
    return dict(kind=kind, div=div, hmin=hmin, ok=span.h >= hmin)


# --------------------------------------------------------------------------
# Design (port of the DESIGN sheet)
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Override:
    """Yellow cells of the DESIGN sheet (None = use the suggested value)."""
    db_top: float | None = None
    db_bot: float | None = None
    layers_top: int = 1
    layers_bot: int = 1
    n_top: int | None = None
    n_bot: int | None = None
    stir: float | None = None
    s: float | None = None     # stirrup spacing (cm)


def _bar_area(d_mm: float) -> float:
    return math.pi * d_mm ** 2 / 400   # cm^2


def _required(st: Settings, b, d, dp_comp, Mu):
    """Singly / doubly design for one moment sign (DESIGN!P:W)."""
    fc, fy, phib, rmax = st.fc, st.fy, st.phib, st.rhomax
    Mn1 = rmax * b * d * fy * (d - rmax * d * fy / (1.7 * fc)) / 1e5
    phiMn1 = phib * Mn1
    out = dict(phiMn1=phiMn1, type="-", Rn=None, rho=None, M2=0.0, fsp=None, As=0.0, Asp=0.0,
               ng=False, need_h=False)
    if Mu <= 0.0001:
        return out
    if Mu <= phiMn1:
        Rn = Mu * 1e5 / (phib * b * d ** 2)
        disc = 1 - 2 * Rn / (0.85 * fc)
        out.update(type="Singly", Rn=Rn)
        if disc < 0:
            out.update(rho=None, ng=True, As=0.0)
            return out
        rho = 0.85 * fc / fy * (1 - math.sqrt(disc))
        out.update(rho=rho, As=max(rho, st.rhomin) * b * d)
        return out
    M2 = Mu / phib - Mn1
    cmax = rmax * d * fy / (0.85 * fc * st.beta1)
    fsp = min(fy, st.es_cu * (cmax - dp_comp) / cmax)
    As = rmax * b * d + M2 * 1e5 / (fy * (d - dp_comp))
    out.update(type="Doubly", rho=rmax, M2=M2, fsp=fsp, As=As)
    if fsp <= 0.85 * fc:
        out.update(need_h=True, Asp=0.0)
    else:
        out["Asp"] = M2 * 1e5 / ((fsp - 0.85 * fc) * (d - dp_comp))
    return out


def _capacity(st: Settings, b, d, dp, As_t, As_c):
    """Strain-compatibility check of the provided bars (DESIGN!AR:AZ)."""
    fc, fy, b1, E = st.fc, st.fy, st.beta1, st.es_cu
    c_y = (As_t * fy - As_c * (fy - 0.85 * fc)) / (0.85 * fc * b1 * b)
    yield_c = c_y > dp and E * (c_y - dp) / c_y >= fy
    B = As_c * (E - 0.85 * fc) - As_t * fy
    c_e = (-B + math.sqrt(B * B + 4 * 0.85 * fc * b1 * b * E * As_c * dp)) / (2 * 0.85 * fc * b1 * b)
    c = c_y if yield_c else c_e
    fsp = fy if yield_c else E * (c - dp) / c
    eps_t = 0.003 * (d - c) / c
    ey = fy / st.Es
    phi = st.phib if st.is99 else max(0.65, min(0.9, 0.65 + 0.25 * (eps_t - ey) / (0.005 - ey)))
    phiMn = phi * (0.85 * fc * b1 * c * b * (d - b1 * c / 2) + As_c * (fsp - 0.85 * fc) * (d - dp)) / 1e5
    if st.is99:
        duct = (As_t - As_c * fsp / fy) / (b * d) <= st.rhomax
    else:
        duct = eps_t >= 0.004
    return dict(c_yield=c_y, yield_c=yield_c, c_elastic=c_e, c=c, fsp=fsp, eps_t=eps_t, phi=phi,
                phiMn=phiMn, ductile=duct)


def design_moments(sm: dict, pos: str) -> tuple[float, float]:
    """(Mu-, Mu+) for a location, from the analysis summary."""
    if pos == "ซ้าย":
        return max(0.0, -sm["ML_min"]), max(0.0, sm["ML_max"])
    if pos == "กลาง":
        return max(0.0, -sm["Mmid_min"]), max(0.0, sm["Mpos_max"])
    return max(0.0, -sm["MR_min"]), max(0.0, sm["MR_max"])


def design_shear_vu(sm: dict, pos: str) -> float:
    return sm["Vu_left"] if pos == "ซ้าย" else (sm["Vu_mid"] if pos == "กลาง" else sm["Vu_right"])


def design_location(st: Settings, sm: dict, pos: str, ov: Override = Override()) -> dict:
    b, h = sm["b"], sm["h"]
    Mneg, Mpos = design_moments(sm, pos)
    dbt = ov.db_top or st.db
    dbb = ov.db_bot or st.db
    stir_mm = st.stir
    d_neg = h - st.cover - st.stir / 10 - (dbt / 10 + 1.25 if ov.layers_top == 2 else dbt / 20)
    d_pos = h - st.cover - st.stir / 10 - (dbb / 10 + 1.25 if ov.layers_bot == 2 else dbb / 20)
    dp_top = st.cover + stir_mm / 10 + dbt / 20
    dp_bot = st.cover + stir_mm / 10 + dbb / 20

    neg = _required(st, b, d_neg, dp_bot, Mneg)     # top in tension, bottom compression
    pos_ = _required(st, b, d_pos, dp_top, Mpos)    # bottom in tension, top compression
    As_top_req = max(neg["As"], pos_["Asp"])
    As_bot_req = max(neg["Asp"], pos_["As"])
    n_top = ov.n_top if ov.n_top else max(2, math.ceil(As_top_req / _bar_area(dbt) - 1e-9))
    n_bot = ov.n_bot if ov.n_bot else max(2, math.ceil(As_bot_req / _bar_area(dbb) - 1e-9))
    As_top = n_top * _bar_area(dbt)
    As_bot = n_bot * _bar_area(dbb)
    bad = neg["need_h"] or pos_["need_h"] or neg["ng"] or pos_["ng"]
    as_ok = (not bad) and As_top >= As_top_req - 0.005 and As_bot >= As_bot_req - 0.005

    def per_layer(db):
        s = max(2.5, db / 10)
        return int((b - 2 * st.cover - 2 * st.stir / 10 + s) // (db / 10 + s))

    fit_top, fit_bot = per_layer(dbt), per_layer(dbb)
    fit_ok = n_top <= fit_top * ov.layers_top and n_bot <= fit_bot * ov.layers_bot

    cap_neg = _capacity(st, b, d_neg, dp_bot, As_top, As_bot)
    cap_pos = _capacity(st, b, d_pos, dp_top, As_bot, As_top)
    strength_ok = (cap_neg["phiMn"] >= Mneg - 1e-4 and cap_pos["phiMn"] >= Mpos - 1e-4
                   and cap_neg["ductile"] and cap_pos["ductile"])
    flex_ok = as_ok and strength_ok and fit_ok
    return dict(
        pos=pos, b=b, h=h, Mneg=Mneg, Mpos=Mpos, db_top=dbt, db_bot=dbb,
        layers_top=ov.layers_top, layers_bot=ov.layers_bot,
        d_neg=d_neg, d_pos=d_pos, dp_top=dp_top, dp_bot=dp_bot, neg=neg, pos_=pos_,
        As_top_req=As_top_req, As_bot_req=As_bot_req, n_top=n_top, n_bot=n_bot,
        As_top=As_top, As_bot=As_bot, top_txt=f"{n_top}-DB{dbt:g}", bot_txt=f"{n_bot}-DB{dbb:g}",
        as_ok=as_ok, fit_top=fit_top, fit_bot=fit_bot, fit_ok=fit_ok,
        cap_neg=cap_neg, cap_pos=cap_pos, strength_ok=strength_ok, flex_ok=flex_ok,
    )


def design_shear(st: Settings, sm: dict, pos: str, flex: dict, ov: Override = Override()) -> dict:
    """Stirrup design (DESIGN!A42:U71); units cm, t."""
    b = sm["b"]
    d = min(flex["d_neg"], flex["d_pos"])
    Vu = design_shear_vu(sm, pos)
    sfc = math.sqrt(st.fc)
    phiVc = st.phiv * 0.53 * sfc * b * d / 1000
    stir = ov.stir or st.stir
    Av = st.legs * math.pi * stir ** 2 / 400
    phiVs_req = max(0.0, Vu - phiVc)
    phiVs_max = st.phiv * 2.1 * sfc * b * d / 1000
    section_ok = phiVs_req <= phiVs_max
    phiVs_tight = st.phiv * 1.1 * sfc * b * d / 1000
    s_str = st.phiv * Av * st.fys * d / (phiVs_req * 1000) if phiVs_req > 0 else None
    s_max = min(d / 4, 30) if phiVs_req > phiVs_tight else min(d / 2, 60)
    s_avmin = Av * st.fys / (3.5 * b)
    s_ctrl = min(s_str if s_str is not None else 1e9, s_max, s_avmin)
    s_sugg = max(5.0, math.floor(s_ctrl / 2.5 + 1e-9) * 2.5)
    s_use = ov.s or s_sugg
    ok = s_use <= s_ctrl + 0.001 and section_ok
    if Vu <= phiVc / 2:
        note = "Vu≤φVc/2 ใส่ขั้นต่ำ"
    elif Vu <= phiVc:
        note = "φVc/2<Vu≤φVc ใช้ Av min"
    else:
        note = "คำนวณเหล็กปลอก"
    label = f"{'DB' if stir >= 10 else 'RB'}{stir:g} @ {s_use / 100:.3f} m"
    return dict(pos=pos, b=b, d=d, Vu=Vu, phiVc=phiVc, stir=stir, Av=Av, phiVs_req=phiVs_req,
                phiVs_max=phiVs_max, section_ok=section_ok, phiVs_tight=phiVs_tight, s_str=s_str,
                s_max=s_max, s_avmin=s_avmin, s_ctrl=s_ctrl, s_sugg=s_sugg, s_use=s_use, ok=ok,
                note=note, label=label)


def design_all(st: Settings, summary: list[dict], overrides: dict | None = None) -> list[dict]:
    """Design every span/location.  overrides: {(span, pos): Override}."""
    overrides = overrides or {}
    rows = []
    for sm in summary:
        for pos in POSITIONS:
            ov = overrides.get((sm["span"], pos), Override())
            fl = design_location(st, sm, pos, ov)
            sh = design_shear(st, sm, pos, fl, ov)
            rows.append(dict(span=sm["span"], pos=pos, flex=fl, shear=sh,
                             ok=fl["flex_ok"] and sh["ok"]))
    return rows


# --------------------------------------------------------------------------
# Load helper (INPUT!K45:Q53)
# --------------------------------------------------------------------------
HELPER_KINDS = ["พื้นทางเดียว (wS/2)", "สองทาง-ด้านสั้น (wS/3)", "สองทาง-ด้านยาว",
                "ผนัง (w×สูง)", "กว้างรับน้ำหนัก (w×B)"]


def helper_load(kind: str, S: float, L: float | None, w: float) -> float | None:
    """Slab / wall load (kg/m^2) -> line load on the beam (t/m)."""
    if not kind or S is None or w is None:
        return None
    if kind == "พื้นทางเดียว (wS/2)":
        v = w * S / 2
    elif kind == "สองทาง-ด้านสั้น (wS/3)":
        v = w * S / 3
    elif kind == "สองทาง-ด้านยาว":
        if not L:
            return None
        v = w * S / 3 * (3 - (S / L) ** 2) / 2
    else:
        v = w * S
    return v / 1000


# --------------------------------------------------------------------------
# Example data (as shipped in the workbook)
# --------------------------------------------------------------------------
EXAMPLE_SPANS = [Span(1.5, 20, 50), Span(4, 25, 50), Span(6, 25, 40), Span(4, 25, 50)]
EXAMPLE_NODES = [
    Node("Free"),
    Node("Column", 20, (20, 20, 3), (20, 20, 3)),
    Node("Column", 25, (25, 25, 3), (25, 25, 3)),
    Node("Column", 25, (25, 25, 3), (25, 25, 3)),
    Node("Pin", 20),
]
EXAMPLE_LOADS = [
    Load(1, "DL", "UDL", 0.65), Load(1, "LL", "UDL", 0.15),
    Load(2, "DL", "UDL", 1.2), Load(2, "LL", "UDL", 0.4),
    Load(3, "DL", "UDL", 3.0), Load(3, "LL", "UDL", 1.2),
    Load(3, "DL", "Point", 3.0, a=3.0), Load(3, "LL", "Point", 1.2, a=3.0),
    Load(4, "DL", "UDL", 0.8),
    Load(4, "DL", "UDL", 0.0, 0.5, 0.0, 2.0), Load(4, "DL", "UDL", 0.5, 0.0, 2.0, 4.0),
    Load(4, "LL", "UDL", 0.0, 0.3, 0.0, 2.0), Load(4, "LL", "UDL", 0.3, 0.0, 2.0, 4.0),
]
EXAMPLE_OVERRIDES = {
    (3, "ซ้าย"): Override(layers_top=2, stir=12),
    (3, "กลาง"): Override(layers_bot=2),
    (3, "ขวา"): Override(layers_top=2, stir=12),
}


# --------------------------------------------------------------------------
# Import from the original workbook (.xlsm / .xlsx)
# --------------------------------------------------------------------------
def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _yes(v) -> bool:
    return str(v).strip().upper() == "YES"


def read_workbook(fp) -> dict:
    """Read INPUT (and the overridden yellow cells of DESIGN) from the workbook.

    Returns dict(settings, spans, nodes, loads, overrides, project, helper).
    """
    import openpyxl

    wb = openpyxl.load_workbook(fp, data_only=False, keep_vba=False)
    if "INPUT" not in wb.sheetnames:
        raise InputError("ไม่พบแผ่นงาน INPUT ในไฟล์")
    ws = wb["INPUT"]
    g = lambda a: ws[a].value
    code = str(g("C15") or CODE99)
    st = Settings(
        fc=_num(g("C10")) or 280.0, fy=_num(g("C11")) or 4000.0, fys=_num(g("C12")) or 2400.0,
        Es=_num(g("C13")) or 2_040_000.0, gc=_num(g("C14")) or 2.4,
        code=CODE19 if code.startswith("ACI 318-19") else CODE99,
        pattern=_yes(g("C16")), self_weight=_yes(g("C17")),
        cover=_num(g("C21")) or 3.0, stir=_num(g("C22")) or 9.0, legs=int(_num(g("C23")) or 2),
        db=_num(g("C24")) or 20.0, all_simple=_yes(g("C28")),
    )
    n = int(_num(g("C27")) or 1)
    n = max(1, min(MAXS, n))
    spans = []
    for r in range(30, 30 + n):
        spans.append(Span(_num(g(f"C{r}")) or 0.0, _num(g(f"D{r}")) or 0.0, _num(g(f"E{r}")) or 0.0,
                          str(g(f"B{r}") or f"B{r - 29}")))
    nodes = []
    for r in range(30, 31 + n):
        kind = str(g(f"L{r}") or "Free").strip().capitalize()
        nn = lambda c: _num(g(f"{c}{r}")) or 0.0
        nodes.append(Node(kind if kind in SUPPORT_TYPES else "Free", nn("M"),
                          (nn("N"), nn("O"), nn("P")), (nn("Q"), nn("R"), nn("S")),
                          str(g(f"T{r}") or "Yes").strip().upper() != "NO"))
    loads = []
    for r in range(47, 107):
        sp, w1 = _num(g(f"B{r}")), _num(g(f"E{r}"))
        if sp is None or w1 is None:
            continue
        loads.append(Load(int(sp), str(g(f"C{r}") or "DL").strip().upper(),
                          "Point" if str(g(f"D{r}") or "").strip().upper() == "POINT" else "UDL",
                          w1, _num(g(f"F{r}")), _num(g(f"G{r}")), _num(g(f"H{r}")),))
    helper = []
    for r in range(46, 54):
        if g(f"L{r}") or g(f"M{r}"):
            helper.append(dict(item=g(f"L{r}") or "", kind=g(f"M{r}") or "", S=_num(g(f"N{r}")),
                               L=_num(g(f"O{r}")), w=_num(g(f"P{r}"))))
    project = {k: g(a) for k, a in dict(project="B4", location="B5", member="B6", type="B7", date="G4",
                                        job="G5", eng="H6", eng_lic="I6", chk="H7", chk_lic="I7").items()}

    overrides = {}
    if "DESIGN" in wb.sheetnames:
        d = wb["DESIGN"]
        const = lambda v: _num(v) if not (isinstance(v, str) and v.startswith("=")) else None
        for i in range(30):
            span, pos = i // 3 + 1, POSITIONS[i % 3]
            if span > n:
                break
            r, rs = 7 + i, 42 + i
            ov = Override(
                db_top=const(d[f"H{r}"].value), db_bot=const(d[f"I{r}"].value),
                layers_top=int(const(d[f"J{r}"].value) or 1), layers_bot=int(const(d[f"K{r}"].value) or 1),
                n_top=int(const(d[f"AH{r}"].value)) if const(d[f"AH{r}"].value) else None,
                n_bot=int(const(d[f"AI{r}"].value)) if const(d[f"AI{r}"].value) else None,
                stir=const(d[f"H{rs}"].value), s=const(d[f"R{rs}"].value),
            )
            if ov != Override():
                overrides[(span, pos)] = ov
    return dict(settings=st, spans=spans, nodes=nodes, loads=loads, overrides=overrides,
                project=project, helper=helper)
