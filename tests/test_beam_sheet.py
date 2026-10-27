"""Design sheet (Singly & Doubly) against the user's Excel forms.

Beam-Singly : B_FL.2_A.6-7.xlsx  (HB4: Mu 5,200 kg-m, Vu 6,000 kg, 20x50, fc' 320, 2-DB16)
Beam-Doubly : B_FL.2_J.1-2_db.xlsx (Mu 17,600 kg-m, Vu 11,800 kg, 20x45, fc' 280, 6-DB20 / 4-DB20)
"""
import pytest

import math
import subprocess
import shutil
from dataclasses import replace

from beam_sheet import Bar, SheetInput, compute, layer_table, suggest
from beam_sheet_report import render_document, render_sheet
from beam_sheet_xlsx import build_workbook


def singly():
    return SheetInput(fc=320, fy=4000, Mu=5200, Vu=6000, Vu_mid=2550, span=3.45, B=20, H=50, cover=4,
                      db=16, bars=[Bar("t", 1, 2, 16), Bar("c", 1, 2, 16)], ds=9, legs=2)


def doubly():
    return SheetInput(fc=280, fy=4000, Mu=17600, Vu=11800, Vu_mid=7615, span=6, B=20, H=45, cover=4,
                      db=20, ds=9, legs=2,
                      bars=[Bar("t", 1, 2, 20), Bar("t", 2, 2, 20), Bar("t", 3, 2, 20),
                            Bar("c", 1, 2, 20), Bar("c", 1, 2, 20)])   # form: both compression rows at d' = 5.9


@pytest.mark.parametrize("key,ref", [
    ("d", 44.3), ("beta1", 0.8214285714285714), ("rho_b", 0.033779220779220784),
    ("rho_max", 0.025334415584415588), ("rho_min", 0.003577708763999664), ("Ru", 59.14330285461293),
    ("d_sugg", 22.10103601353327), ("As_sugg", 7.4655577495065), ("Rn", 14.720527946073041),
    ("rho_req", 0.003785499691786775), ("As_req", 3.3539527269230827), ("rho", 0.004538644014215503),
    ("a", 2.9567930857315705), ("phiMn", 6199.051845236805),
])
def test_singly_form(key, ref):
    r = compute(singly())
    assert r["kind"] == "Singly"
    assert r[key] == pytest.approx(ref, rel=1e-9)


def test_singly_form_shear():
    r = compute(singly())
    z, m = r["sh_2h"], r["sh_mid"]
    assert z["phiVc"] == pytest.approx(7140.087045945589) and z["lim11"] == pytest.approx(14819.048585924811)
    assert z["lim21"] == pytest.approx(28290.91093676554)
    assert (z["case"], z["s"]) == (2, 10.0)          # EQ: min(d/4, 8db, 24ds, 30) = 11.075 -> 10
    assert (m["case"], m["s"]) == (1, 20.0)
    assert r["ok"]


@pytest.mark.parametrize("key,ref", [
    ("d", 34.1), ("dp", 5.9), ("rho_b", 0.030584881422924905), ("rho1", 0.015292440711462452),
    ("rho_max", 0.02293866106719368), ("rho_cy", 0.025260915453992144), ("P", 0.18635250540325002),
    ("c", 10.405232126177903), ("fs_c", 2649.822731281695), ("Ru1", 53.2853823704772),
    ("Mu1", 11152.939585358628), ("Mu2", 6447.060414641372), ("Asc_req", 9.586350435230976),
    ("As_req", 16.779976888544194), ("duct_val", 0.015432393395350361), ("a", 8.844447307251215),
    ("phiMn", 19695.992899320416),
])
def test_doubly_form(key, ref):
    r = compute(doubly())
    assert r["kind"] == "Doubly"
    assert r[key] == pytest.approx(ref, rel=1e-9)


def test_doubly_form_shear():
    r = compute(doubly())
    z, m = r["sh_2h"], r["sh_mid"]
    assert z["phiVc"] == pytest.approx(5141.125264247119) and z["case"] == 3
    assert z["s_ctrl"] == pytest.approx(8.525)        # EQ limit d/4
    assert m["case"] == 3 and m["s_ctrl"] == pytest.approx(17.05)
    assert r["flex_ok"] and r["sh_2h"]["ok"] and r["sh_mid"]["ok"]
    # the user's extra checks flag the form's section: 6-DB20 in 20x45 = 2.09 % > 2 %, 4-DB20 in one row
    assert r["pct"] == pytest.approx(6 * math.pi / (20 * 45)) and not r["pct_ok"]
    assert not r["clear_ok"] and not r["ok"]


@pytest.mark.parametrize("Mu", [500, 3000, 8000, 15000, 22000, 30000])
def test_suggestion_passes(Mu):
    r = compute(suggest(SheetInput(fc=240, fy=4000, Mu=Mu, Vu=9000, Vu_mid=6000, B=25, H=50, db=20)))
    assert r["ok"], r["kind"]


def test_forced_doubly_when_rounded_singly_fails():
    """rho required just below rho_max: the bars needed fail the singly check -> designed as doubly."""
    s = suggest(SheetInput(fc=280, fy=4000, fy_rb=2400, Mu=15600, Vu=15000, Vu_mid=10000, span=6, B=25,
                           H=40, cover=3, db=20, ds=12, legs=2))
    r = compute(s)
    assert r["kind"] == "Doubly" and s.mode == "doubly" and r["auto_kind"] == "Singly" and r["ok"]


def test_render_page():
    page = render_sheet(compute(doubly()), dict(project="P", member="B1", page="1/1"))
    doc = render_document([page, render_sheet(compute(singly()), {})])
    assert doc.count('class="sheet"') == 2 and "Doubly" in doc and "Singly" in doc and "19,696" in doc


def test_clear_spacing_formula():
    """(B - 2 cover - stirrup x legs - n db)/(n - 1), e.g. (20 - 8 - 0.9x2 - 2x1.6)/1 = 7.0 cm."""
    lt = layer_table(singly(), "t")
    assert lt[0]["clear"] == pytest.approx(7.0)
    three = replace(singly(), bars=[Bar("t", 1, 3, 16)], legs=3)
    assert layer_table(three, "t")[0]["clear"] == pytest.approx((20 - 8 - 0.9 * 3 - 4.8) / 2)
    assert not compute(three)["clear_ok"]


def test_mixed_sizes_and_layers():
    inp = replace(singly(), bars=[Bar("t", 1, 2, 20), Bar("t", 1, 1, 16), Bar("t", 2, 2, 12), Bar("c", 1, 2, 12)])
    lt = layer_table(inp, "t")
    assert [l["n"] for l in lt] == [3, 2] and lt[0]["dmax"] == 20
    assert lt[1]["y"] == pytest.approx(4 + 0.9 + 2.0 + 3 + 0.6)
    r = compute(inp)
    assert r["d"] == pytest.approx(50 - (lt[0]["As"] * lt[0]["y"] + lt[1]["As"] * lt[1]["y"]) / (lt[0]["As"] + lt[1]["As"]))


def test_mode_switch_back_to_singly():
    """A doubly section can be designed as singly on request; it then fails the singly checks honestly."""
    r = compute(replace(doubly(), mode="singly"))
    assert r["kind"] == "Singly" and r["auto_kind"] == "Doubly" and r["chosen"] and not r["duct_ok"]
    s = suggest(replace(doubly(), mode="singly"))
    assert s.mode == "singly" and compute(s)["kind"] == "Singly"
    s2 = suggest(replace(singly(), mode="doubly"))
    r2 = compute(s2)
    assert r2["kind"] == "Doubly" and r2["flex_ok"]


def test_suggestion_respects_clear_and_adds_legs():
    s = suggest(SheetInput(fc=240, Mu=12000, Vu=38000, Vu_mid=20000, B=30, H=60, db=20, ds=9, legs=2))
    r = compute(s)
    assert r["clear_ok"] and r["ok"]
    assert min(r["sh_2h"]["s"], r["sh_mid"]["s"]) >= 7.5 and (s.ds > 9 or s.legs > 2)


def test_xlsx_has_formulas():
    from openpyxl import load_workbook
    import io
    data = build_workbook([("S1-กลาง", doubly(), {"member": "B1"}), ("S1-ซ้าย", singly(), {})])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["สรุป", "S1-กลาง", "S1-ซ้าย"]
    ws = wb["S1-กลาง"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 80 and any("SQRT" in f for f in formulas)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_xlsx_recalculates_like_python(tmp_path):
    """Recalculate the workbook in LibreOffice and compare with compute()."""
    from openpyxl import load_workbook
    cases = [singly(), doubly(), replace(doubly(), mode="singly"),
             replace(singly(), bars=[Bar("t", 1, 2, 20), Bar("t", 1, 1, 16), Bar("t", 2, 2, 12), Bar("c", 1, 2, 12)],
                     legs=3, ds=12, eq=False)]
    src = tmp_path / "in.xlsx"
    src.write_bytes(build_workbook([(f"c{i}", c, {}) for i, c in enumerate(cases)]))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc Office Open XML", "--outdir",
                    str(tmp_path / "out"), str(src)], check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmp_path / "out" / "in.xlsx", data_only=True)
    for row, inp in zip(wb["สรุป"].iter_rows(min_row=2, values_only=True), cases):
        r = compute(inp)
        assert row[1] == r["kind"]
        for got, ref in zip(row[2:6], [inp.Mu, r["As"], r["Asc"], r["d"]]):
            assert got == pytest.approx(ref, rel=1e-9)
        assert row[6] == pytest.approx(r["pct"], rel=1e-9)
        assert row[8] == pytest.approx(r["phiMn"], rel=1e-9)
        assert (row[9], row[10]) == (r["sh_2h"]["s"], r["sh_mid"]["s"])
        assert row[11] == ("ผ่าน" if r["ok"] else "ไม่ผ่าน")
